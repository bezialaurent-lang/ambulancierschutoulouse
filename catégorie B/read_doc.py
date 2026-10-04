import docx

doc = docx.Document('Analyse_Revalorisation_Diplome_Ambulancier_Dossier_Complet.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print()

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        print(f"[{i:03d}] Style={para.style.name!r:<30} | {text[:100]}")
